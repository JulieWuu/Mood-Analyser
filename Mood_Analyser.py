import pandas as pd
import numpy as np
import re

from openpyxl import load_workbook
from matplotlib import pyplot as plt
from Excel_to_Json import get_next_date
from TinyDB_of_Mood import DONT_CARE_LIST

def lexify(comment):
    cleaned = re.sub(r'[^\w\s]', '', comment.lower())
    return cleaned.split(" ")

def read_excel(path):
    work_sheet = load_workbook(filename=path, read_only=True).active
    data_structure = {
        "mood": [],
        "date": [],
        "month": [],
        "day": [],
        "week": [],
        "comment": [],
        "lexicons": [],
    }

    date_pair = tuple((1, 1))
    for i in range(3, work_sheet.max_row + 1, 2):
        for j in range(1, work_sheet.max_column + 1):
            if i == 3 and j <= 3: continue
            cell_obj = work_sheet.cell(row=i, column=j)
            if cell_obj.value is None: continue

            data_structure["month"].append(date_pair[0])
            data_structure["date"].append(date_pair[1])
            data_structure["day"].append((j+5) % 7 + 1)
            data_structure["week"].append(i // 2)
            data_structure["comment"].append(cell_obj.value)
            data_structure["lexicons"].append(lexify(cell_obj.value))

            colour = cell_obj.font.color
            mood = 0.
            if colour.type == 'rgb':
                if colour.rgb == 'FF0070C0':
                    mood = 1.
                elif colour.rgb == 'FF00B050':
                    mood = 1 / 2
            elif colour.type == 'theme' and colour.theme == 6:
                mood = 1 / 3
            data_structure["mood"].append(mood)
            date_pair = get_next_date(date_pair)
            if date_pair == (1, 1): break

    return pd.DataFrame(data_structure)

def plot_mood(df, axis=""):
    if not axis: average_mood = df['mood']
    else: average_mood = df.groupby(axis)['mood'].mean()

    fig, ax = plt.subplots(figsize=(10, 5))
    ax.spines['top'].set_visible(False)
    ax.set_title(f"Average mood by {axis}" if axis else "Mood score per day")
    ax.set_ylim(np.floor(average_mood.min() * 2.5) * .4,1.01)
    ax.plot(average_mood)
    plt.savefig(f"Mood_by_{axis}.png")
    plt.show()

def plot_words_by_count(df, count_boundary=(10, 400)):
    count_map = {}
    for idx, words in df["lexicons"].items():
        for w in words:
            if w not in DONT_CARE_LIST:
                if w not in count_map:
                    count_map[w] = [1, 0, 0, 0]
                else: count_map[w][0] += 1
                count_map[w][int(1 / df["mood"][idx])] += 1

    filtered_map = {k : v for k, v in count_map.items() if count_boundary[0] <= v[0] <= count_boundary[1]}

    fig, [ax1, ax2] = plt.subplots(nrows=1, ncols=2, figsize=(15, 8))

    rank_by_count = {k: v for k, v in sorted(filtered_map.items(), key=lambda item: item[1][0])}
    x1 = list(rank_by_count.keys())
    ax1.set_title("Word by number of appearances")
    ax1.barh(x1, [c[3] for c in rank_by_count.values()], color='orange')
    ax1.barh(x1, [c[2] for c in rank_by_count.values()], left=[c[3] for c in rank_by_count.values()], color='green')
    ax1.barh(x1, [c[1] for c in rank_by_count.values()], left=[c[2]+c[3] for c in rank_by_count.values()], color='blue')

    rank_by_happiness = {k: v for k, v in sorted(filtered_map.items(), key=lambda item: (item[1][1] / item[1][0]))}
    x2 = list(rank_by_happiness.keys())
    ax2.set_title("Word by percentage of happiness")
    ax2.barh(x2, [c[3] / c[0] for c in rank_by_happiness.values()], color='orange')
    ax2.barh(x2, [c[2] / c[0] for c in rank_by_happiness.values()], left=[c[3] / c[0] for c in rank_by_happiness.values()], color='green')
    ax2.barh(x2, [c[1] / c[0] for c in rank_by_happiness.values()], left=[(c[2] + c[3]) / c[0] for c in rank_by_happiness.values()], color='blue')

    plt.savefig("Word_plot.png")
    plt.show()

if __name__ == "__main__":
    mood_df = read_excel("EVERYDAYlog.xlsx")
    plot_mood(mood_df, 'day')